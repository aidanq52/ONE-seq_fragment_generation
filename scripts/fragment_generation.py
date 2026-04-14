# scripts/fragment_generation.py
import pandas as pd
import os
import re
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

# Constants
LEFT_PBS = "GACGTTCTCACAGCAATTCGTACAGTCGACGTCGATTCGTGT"
PROTO_CONST = "TTGACATTCTGCAATTA"
PAM_COST = "AGTATGTATGCTTCGCGCAGTGCGACTTCGCAGCGCATCACTTCA"
RIGHT_PBS = "AGAGCTGCGAGTCTTACAGCATTGC"

def generate_fragments(
    input_file=None,
    barcode_file="barcode_list.txt",
    output_file=None
):
    if input_file is None:
        dedup_file = "intermediate_files/5a_pre_barcode_plus_variants.txt"
        fallback_file = "intermediate_files/3a_cleaned_file_with_sequences.txt"
        input_file = dedup_file if os.path.exists(dedup_file) else fallback_file

    if output_file is None:
        user_filename = input("Enter a name for the final oligo output file (without extension): ").strip()
        if not user_filename:
            raise ValueError("❌ Output filename cannot be empty.")
        if not user_filename.endswith(".txt"):
            user_filename += ".txt"
        output_file = os.path.join(os.getcwd(), user_filename)


    if not os.path.exists(input_file):
        raise FileNotFoundError(f"❌ Input file not found: {input_file}")
    
    if not os.path.exists(barcode_file):
        raise FileNotFoundError(f"❌ Barcode file not found: {barcode_file}")

    # Read data
    df = pd.read_csv(input_file, sep='\t')

    # Read and reverse barcodes
    with open(barcode_file, 'r') as f:
        barcodes = [line.strip() for line in f if line.strip()]
    barcodes.reverse()

    if len(barcodes) < len(df):
        raise ValueError("❌ Not enough barcodes for the number of rows in the input file.")

    # Assign barcodes
    df['Left_PBS'] = LEFT_PBS
    df['barcode'] = barcodes[:len(df)]
    df['proto_const'] = PROTO_CONST
    df['PAM_cost'] = PAM_COST
    df['barcode_2'] = df['barcode']
    df['Right_PBS'] = RIGHT_PBS

    # Reorder columns
    fetched_seq = df.pop('fetched_sequence')
    df.insert(df.columns.get_loc('proto_const') + 1, 'fetched_sequence', fetched_seq)

    # Generate oligo
    df['oligo'] = (
        df['Left_PBS'].astype(str) +
        df['barcode'].astype(str) +
        df['proto_const'].astype(str) +
        df['fetched_sequence'].astype(str) +
        df['PAM_cost'].astype(str) +
        df['barcode_2'].astype(str) +
        df['Right_PBS'].astype(str)
    )

    df.to_csv(output_file, sep='\t', index=False)
    print(f"Final annotated oligo file saved to: {output_file}")

    # Write Excel workbook with two sheets: library + multi-site report
    xlsx_file = output_file.replace(".txt", ".xlsx")
    multi_site_file = "intermediate_files/4c_multi_site_report.csv"
    with pd.ExcelWriter(xlsx_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Library", index=False)
        if os.path.exists(multi_site_file):
            multi_df = pd.read_csv(multi_site_file)
            multi_df.to_excel(writer, sheet_name="Multi-Site Report", index=False)

            last_col = get_column_letter(len(multi_df.columns))

            # Walk the Multi-Site Report rows to find each reference group's row span
            group_spans = {}   # frag_numb -> (start_excel_row, end_excel_row)
            current_ref = None
            current_start = None
            for idx, row in multi_df.iterrows():
                excel_row = idx + 2  # +1 for header, +1 for 1-based Excel rows
                if row["role"] == "reference":
                    current_ref = row["frag_numb"]
                    current_start = excel_row
                    group_spans[current_ref] = (excel_row, excel_row)
                elif current_ref is not None:
                    group_spans[current_ref] = (current_start, excel_row)

            # Create a named range per group so the hyperlink selects the whole block
            wb = writer.book
            for frag_numb, (start_row, end_row) in group_spans.items():
                safe_name = "REF_" + re.sub(r"[^A-Za-z0-9_]", "_", frag_numb)
                range_ref = f"'Multi-Site Report'!$A${start_row}:${last_col}${end_row}"
                dn = DefinedName(name=safe_name, attr_text=range_ref)
                wb.defined_names.add(dn)

            # Add hyperlinks in Library sheet pointing to each group's named range
            lib_sheet = writer.sheets["Library"]
            for excel_row in range(2, len(df) + 2):
                cell = lib_sheet.cell(row=excel_row, column=1)
                if cell.value in group_spans:
                    safe_name = "REF_" + re.sub(r"[^A-Za-z0-9_]", "_", cell.value)
                    cell.hyperlink = f"#{safe_name}"
                    cell.style = "Hyperlink"
        else:
            print(f"Warning: {multi_site_file} not found — Multi-Site Report sheet skipped.")
    print(f"Excel workbook saved to: {xlsx_file}")
